"""Weekly Excel report against the test database (small sample + quick-model risk scores).

openpyxl does not compute formulas, so a tiny evaluator resolves the handful of formula shapes
the report uses (cell references, arithmetic, SUM, AVERAGE, IFERROR, IF, ROUND) and the
results are compared with a pandas cross-check of the same week.
"""

from __future__ import annotations

import re
from datetime import datetime, timedelta

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from conftest import load_scored_db
from slawatch import report
from slawatch.report import KPI_ROWS, SUMMARY_KPI_HEADER_ROW, SUMMARY_TREND_HEADER_ROW

pytestmark = pytest.mark.integration

SHEETS = ["Summary", "By customer", "By service", "Backlog ageing", "At-risk open tickets", "Notes"]


@pytest.fixture(scope="module")
def scored_engine(small_raw_dir, trained):
    engine = load_scored_db(small_raw_dir, trained)
    yield engine
    engine.dispose()


@pytest.fixture(scope="module")
def built(scored_engine, tmp_path_factory):
    out = tmp_path_factory.mktemp("reports") / "synthetic_weekly_kpi_test.xlsx"
    path, data = report.write_report(scored_engine, out=out, top_n=25)
    wb = load_workbook(path)  # formulas kept as strings
    return path, data, wb


# ------------------------------------------------------------------------------------------
# Minimal formula evaluator
# ------------------------------------------------------------------------------------------
_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")
_RANGE = re.compile(r"\$?([A-Z]{1,3})\$?(\d+):\$?([A-Z]{1,3})\$?(\d+)")


class _Sheet:
    def __init__(self, ws):
        self.ws = ws

    def value(self, ref: str):
        v = self.ws[ref.replace("$", "")].value
        if isinstance(v, str) and v.startswith("="):
            return self.eval(v[1:])
        if isinstance(v, datetime):  # Excel serial days, so date arithmetic behaves as in Excel
            return (v - datetime(1899, 12, 30)).total_seconds() / 86400
        return v

    def _range(self, c1, r1, c2, r2):
        out = []
        for r in range(int(r1), int(r2) + 1):
            for c in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
                out.append(self.value(f"{get_column_letter(c)}{r}"))
        return out

    def eval(self, expr: str):
        expr = _RANGE.sub(lambda m: f"__rng({m[1]!r},{m[2]},{m[3]!r},{m[4]})", expr)
        expr = _REF.sub(lambda m: f"__ref({m[1]!r}{m[2]!r})", expr)
        expr = expr.replace("<>", "!=")
        ns = {
            "__rng": self._range,
            "__ref": self.value,
            "SUM": lambda xs: sum(x for x in xs if isinstance(x, int | float)),
            "AVERAGE": _average,
            "ROUND": lambda x, n: round(x, int(n)),
            "IF": lambda c, a, b: a if c else b,
            "IFERROR": _iferror,
        }
        # IFERROR needs lazy evaluation of its first argument
        expr = re.sub(r"IFERROR\((.*),(\"\")\)$", r"IFERROR(lambda: \1, \2)", expr)
        return eval(expr, {"__builtins__": {}, **ns})  # noqa: S307 - test-only, fixed grammar


def _average(xs):
    v = [x for x in xs if isinstance(x, int | float)]
    return sum(v) / len(v) if v else None


def _iferror(thunk, fallback):
    try:
        v = thunk()
    except (ZeroDivisionError, TypeError):
        return fallback
    return fallback if v is None else v


# ------------------------------------------------------------------------------------------
# Tests
# ------------------------------------------------------------------------------------------
def test_sheets_and_labels(built):
    path, data, wb = built
    assert path.name.startswith("synthetic_")
    assert wb.sheetnames == SHEETS
    for name in SHEETS[:-1]:
        ws = wb[name]
        assert "SYNTHETIC" in str(ws["A3"].value).upper()
        assert ws.freeze_panes is not None
    notes = wb["Notes"]
    assert notes["A3"].value == "SYNTHETIC DATA" and notes["A3"].font.bold
    labels = {notes.cell(row=r, column=1).value for r in range(1, 40)}
    assert {"Generated at (UTC)", "Data source", "Model version", "Definitions"} <= labels
    assert "synthetic" in wb.properties.title.lower()
    assert data.model_version == "test-quick"
    assert wb["Summary"]._charts and wb["Backlog ageing"]._charts and wb["By service"]._charts


def test_default_week_is_last_full_week(built, scored_engine):
    _, data, _ = built
    weeks = report.available_weeks(scored_engine)
    assert data.week_ending == weeks[-1] and data.week_ending.weekday() == 6
    assert data.as_of == datetime.combine(data.week_ending + timedelta(days=1), datetime.min.time())
    assert len(data.trend) == report.TREND_WEEKS
    assert pd.Timestamp(data.trend["week_ending"].iloc[-1]).date() == data.week_ending


def test_summary_formulas_match_the_view_and_pandas(built):
    _, data, wb = built
    ws = wb["Summary"]
    sh = _Sheet(ws)
    trend = data.trend.reset_index(drop=True)
    last, prior = trend.iloc[-1], trend.iloc[-2]
    base = trend.iloc[-1 - report.BASELINE_WEEKS : -1]
    first_kpi = SUMMARY_KPI_HEADER_ROW + 1
    col_this, col_prior, col_dprior, col_avg, col_davg = "B", "C", "D", "E", "F"
    for i, (label, key, _fmt, _, _) in enumerate(KPI_ROWS):
        r = first_kpi + i
        assert ws[f"A{r}"].value == label
        for col in (col_this, col_prior, col_dprior, col_avg, col_davg):
            assert str(ws[f"{col}{r}"].value).startswith("="), (label, col)
        this, prev, avg = sh.value(f"B{r}"), sh.value(f"C{r}"), sh.value(f"E{r}")
        if key == "sla_compliance":
            exp_this = 1 - last["breached_tickets"] / last["tickets_resolved"]
            exp_prev = 1 - prior["breached_tickets"] / prior["tickets_resolved"]
            exp_avg = (1 - base["breached_tickets"] / base["tickets_resolved"]).mean()
            assert sh.value(f"D{r}") == pytest.approx((exp_this - exp_prev) * 100)
        else:
            exp_this, exp_prev = float(last[key]), float(prior[key])
            exp_avg = base[key].astype(float).mean()
            assert sh.value(f"D{r}") == pytest.approx(exp_this - exp_prev)
        assert this == pytest.approx(exp_this), label
        assert prev == pytest.approx(exp_prev), label
        assert avg == pytest.approx(exp_avg), label
        assert sh.value(f"F{r}") == pytest.approx(
            (this - avg) * (100 if key == "sla_compliance" else 1)
        )
    # trend table values are the view's numbers, compliance is a formula
    for i, row in trend.iterrows():
        r = SUMMARY_TREND_HEADER_ROW + 1 + i
        assert ws[f"A{r}"].value.date() == pd.Timestamp(row["week_ending"]).date()
        assert ws[f"B{r}"].value == row["tickets_opened"]
        assert ws[f"C{r}"].value == row["tickets_resolved"]
        assert ws[f"D{r}"].value == row["breached_tickets"]
        assert ws[f"E{r}"].value.startswith("=IFERROR(")
        assert sh.value(f"E{r}") == pytest.approx(float(row["sla_compliance_pct"]) / 100, abs=1e-4)


def test_by_customer_and_service_reconcile(built):
    _, data, wb = built
    last = data.trend.iloc[-1]
    ot = data.open_tickets
    for sheet, first_col, n_rows in (
        ("By customer", 4, len(data.by_customer)),
        ("By service", 2, len(data.by_service)),
    ):
        ws = wb[sheet]
        sh = _Sheet(ws)
        total_row = 6 + n_rows
        assert ws.cell(row=total_row, column=1).value == "Total"
        opened = sh.value(f"{get_column_letter(first_col)}{total_row}")
        resolved = sh.value(f"{get_column_letter(first_col + 1)}{total_row}")
        breaches = sh.value(f"{get_column_letter(first_col + 2)}{total_row}")
        assert (opened, resolved, breaches) == (
            last["tickets_opened"],
            last["tickets_resolved"],
            last["breached_tickets"],
        )
        comp = sh.value(f"{get_column_letter(first_col + 3)}{total_row}")
        assert comp == pytest.approx(1 - breaches / resolved)
    ws = wb["By customer"]
    sh = _Sheet(ws)
    n = len(data.by_customer)
    backlog = sh.value(f"K{6 + n}")
    past_due = sh.value(f"L{6 + n}")
    high = sh.value(f"M{6 + n}")
    assert backlog == len(ot) == last["open_backlog"]
    assert past_due == int(ot["past_due"].sum()) == last["backlog_past_due"]
    assert high == int((ot["risk_band"] == "high").sum()) == last["high_risk_open"]
    # first data row compliance / delta are formulas that evaluate to the pandas figures
    row = data.by_customer.iloc[0]
    if row["tickets_resolved"]:
        assert sh.value("G6") == pytest.approx(
            1 - row["breached_tickets"] / row["tickets_resolved"]
        )


def test_backlog_ageing_totals(built):
    _, data, wb = built
    ws = wb["Backlog ageing"]
    sh = _Sheet(ws)
    ot = data.open_tickets
    groups = sorted(ot["assignment_group"].unique())
    total_row = 7 + len(groups)
    assert ws.cell(row=total_row, column=1).value == "Total"
    assert sh.value(f"G{total_row}") == len(ot)
    assert sh.value(f"H{total_row}") == int(ot["past_due"].sum())
    for j, bucket in enumerate(report.AGE_BUCKETS):
        col = get_column_letter(2 + j)
        assert sh.value(f"{col}{total_row}") == int((ot["age_bucket"] == bucket).sum())


def test_at_risk_sheet_is_top_n_by_probability(built):
    _, data, wb = built
    ws = wb["At-risk open tickets"]
    sh = _Sheet(ws)
    n = len(data.at_risk)
    assert 0 < n <= 25
    probs = [ws.cell(row=7 + i, column=14).value for i in range(n)]
    assert probs == sorted(probs, reverse=True)
    assert probs[0] == pytest.approx(float(data.open_tickets["probability"].max()))
    assert ws["B4"].value == data.as_of
    for i in range(n):
        r = 7 + i
        created = ws.cell(row=r, column=9).value
        age = sh.value(f"J{r}")
        assert age == pytest.approx((data.as_of - created).total_seconds() / 3600, abs=0.06)
        expected = ws.cell(row=r, column=12).value
        assert sh.value(f"M{r}") == ("yes" if data.as_of > expected else "no")
        assert sh.value(f"M{r}") == ("yes" if data.at_risk.iloc[i]["past_due"] else "no")


def test_week_ending_validation(scored_engine):
    weeks = report.available_weeks(scored_engine)
    with pytest.raises(SystemExit, match="no full week ends"):
        report.fetch(scored_engine, weeks[-1] + timedelta(days=7))
    with pytest.raises(SystemExit, match="first week"):
        report.fetch(scored_engine, weeks[0])
    with pytest.raises(Exception, match="not a Sunday"):
        report.parse_week_ending("2026-06-27")
    assert report.parse_week_ending("2026-06-28").weekday() == 6
    earlier = report.fetch(scored_engine, weeks[-5], top_n=5)
    assert earlier.week_ending == weeks[-5] and len(earlier.at_risk) <= 5
